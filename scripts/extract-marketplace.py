import json
from pathlib import Path
import openpyxl

source = Path('C:/Users/LENOVO/Downloads/Silabus Mini Class Marketplace.xlsx')
w = openpyxl.load_workbook(source, data_only=True)
s = w['Soal Quiz Batch 3-4']
def color(c):
    return c.rgb[-6:] if c and c.type == 'rgb' else None
items=[]
for numcol,qcol,labelcol,optcol in [(2,3,4,5),(7,8,9,10)]:
    for row in range(3, s.max_row+1):
        n=s.cell(row,numcol)
        if not isinstance(n.value,(int,float)): continue
        number=int(n.value)
        opts=[]
        for r in range(row,row+4):
            label=s.cell(r,labelcol); option=s.cell(r,optcol)
            opts.append({'label':label.value,'text':option.value,'yellow':color(option.fill.fgColor)=='FFFF00' or color(label.fill.fgColor)=='FFFF00','cell':option.coordinate})
        items.append({'number':number,'cell':n.coordinate,'red':color(n.font.color) in ['FF0000','EA4335','CC0000','EA9999'] or color(n.fill.fgColor) in ['FF0000','EA4335','CC0000','EA9999'],'font':color(n.font.color),'fill':color(n.fill.fgColor),'question':s.cell(row,qcol).value,'options':opts})
items.sort(key=lambda q:q['number'])
output={'source':source.name,'sheet':s.title,'questions':items}
Path('docs/quiz-archive').mkdir(parents=True,exist_ok=True)
Path('docs/quiz-archive/marketplace-source.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(items)==20
assert all(len(q['options'])==4 and sum(o['yellow'] for o in q['options'])==1 for q in items)
print(json.dumps({'sheet':s.title,'total':len(items),'excluded':[q['number'] for q in items if q['red']],'included':len([q for q in items if not q['red']])}))
